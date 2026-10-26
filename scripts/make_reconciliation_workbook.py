"""Excel reconciliation workbook: powerbi/reconciliation.xlsx

    python scripts/make_reconciliation_workbook.py
    python scripts/make_reconciliation_workbook.py --gold-db E:/rajasthan-udise-insights/data/gold/udise.duckdb

Run scripts/export_powerbi.py first. Three sources are compared:

    Official      reference/official_totals.json (published UDISE+ figures)
    Pipeline      the gold database (pipeline/validation.py, sql/validation/)
    Power BI CSVs the exported files in powerbi/data/, with the DAX measures of
                  powerbi/MEASURES.md written in SQL (sql/powerbi/check_*.sql)

Sheets:
    Reconciliation   official vs pipeline vs CSVs; differences and result are Excel formulas
    State measures   expected value of every DAX measure, by year (to check Power BI cards)
    District measures the same per district and year; "vs state" columns are formulas
    Checks           automated checks run by this script (all must pass, else exit code 1)
    Notes            definitions, norms and sources

Only state and district aggregates go into the workbook (fine to publish).
"""
import argparse
import sys
from datetime import datetime, timezone
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pipeline import config as C  # noqa: E402
from pipeline.validation import compute_indicators  # noqa: E402
from scripts.export_powerbi import AGG_SQL, INDICATOR_LABELS, OUT_DIR as CSV_DIR, TABLES  # noqa: E402

OUT = C.ROOT / "powerbi" / "reconciliation.xlsx"
CHECK_SQL = C.ROOT / "sql" / "powerbi"

# State / district measures: (column in check_district_measures.sql, DAX measure name, number format)
MEASURES = [
    ("schools_with_students", "Schools with Students", "count"),
    ("enrolment_1_12", "Enrolment 1-12", "count"),
    ("teachers", "Teachers (schools with students)", "count"),
    ("ptr", "PTR", "ratio"),
    ("pct_single_teacher", "Single-Teacher Schools %", "pct"),
    ("avg_infra_index", "Infrastructure Index", "ratio"),
    ("girls_share", "Girls' Share 1-12", "pct"),
    ("girls_retention_gap", "Girls' Retention Gap", "pct"),
    ("pct_has_electricity", "% Schools with Electricity", "pct"),
    ("pct_has_safe_water", "% Schools with Safe Water", "pct"),
    ("pct_has_girls_toilet", "% Schools with Girls' Toilet", "pct"),
    ("teachers_required", "Teachers Required (Norm)", "count"),
    ("teacher_shortfall", "Teacher Shortfall", "count"),
    ("teacher_surplus", "Teacher Surplus", "count"),
    ("net_teacher_gap", "Net Teacher Gap", "signed"),
    ("teacher_shortfall_pct", "Teacher Shortfall %", "pct"),
    ("schools_short_of_teachers", "Schools Short of Teachers", "count"),
    ("classroom_shortfall", "Classroom Shortfall", "count"),
    ("classroom_shortfall_usable", "Classroom Shortfall (Usable Rooms)", "count"),
]
# CSV measure -> mart_district_year column (and the factor the mart is scaled by)
MART_COLUMNS = {"schools_with_students": ("schools", 1), "enrolment_1_12": ("enrolment", 1),
                "teachers": ("teachers", 1), "ptr": ("ptr", 1), "pct_single_teacher": ("pct_single_teacher", 100),
                "avg_infra_index": ("avg_infra_index", 1), "girls_share": ("girls_share", 1),
                "girls_retention_gap": ("girls_retention_gap", 1),
                "pct_has_electricity": ("pct_has_electricity", 100),
                "pct_has_safe_water": ("pct_has_safe_water", 100),
                "pct_has_girls_toilet": ("pct_has_girls_toilet", 100)}
ADDITIVE = ["schools_with_students", "enrolment_1_12", "teachers", "teachers_required", "teacher_shortfall",
            "teacher_surplus", "schools_short_of_teachers", "classroom_shortfall", "classroom_shortfall_usable"]

# Indian digit grouping, as in UDISE+ reports: 1,63,64,187
INDIAN = "[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0"
FORMATS = {"count": INDIAN, "signed": "#,##0;-#,##0;0", "ratio": "0.00", "pct": "0.0%"}
FONT, BLUE = "Arial", "0000FF"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="BFBFBF")


# --- numbers ------------------------------------------------------------------------

def csv_connection(csv_dir: Path):
    """In-memory DuckDB with one view per exported CSV, named like the Power BI table."""
    con = duckdb.connect()
    for name in [*TABLES, "norms"]:
        path = csv_dir / f"{name}.csv"
        if not path.exists():
            raise FileNotFoundError(f"{path} not found: run python scripts/export_powerbi.py first")
        con.execute(f"create view {name} as select * from read_csv('{path.as_posix()}', header = true)")
    return con


def compute(gold_db: Path, csv_dir: Path) -> dict:
    con = csv_connection(csv_dir)
    try:
        csv_validation = con.execute((CHECK_SQL / "check_validation_measures.sql").read_text(encoding="utf-8")).df()
        measures = con.execute((CHECK_SQL / "check_district_measures.sql").read_text(encoding="utf-8")).df()
        csv_rows = {t: con.execute(f"select count(*) from {t}").fetchone()[0] for t in TABLES}
    finally:
        con.close()
    gold = duckdb.connect(str(gold_db), read_only=True)
    try:
        mart = gold.execute("select * from mart_district_year").df()
        gold_rows = {t: gold.execute(f"select count(*) from ({q or AGG_SQL.read_text(encoding='utf-8')})").fetchone()[0]
                     for t, q in TABLES.items()}
    finally:
        gold.close()
    return {"pipeline": compute_indicators(gold_db), "csv_validation": csv_validation, "measures": measures,
            "mart": mart, "csv_rows": csv_rows, "gold_rows": gold_rows}


def run_checks(d: dict) -> list[tuple[str, bool, str]]:
    checks = []
    for t in TABLES:
        checks.append((f"Rows in {t}.csv = rows in gold", d["csv_rows"][t] == d["gold_rows"][t],
                       f"{d['csv_rows'][t]:,} in CSV, {d['gold_rows'][t]:,} in gold"))
    both = d["pipeline"].merge(d["csv_validation"], on=["year", "indicator"], how="outer",
                               suffixes=("_gold", "_csv"))
    diff = both[both["value_gold"] != both["value_csv"]]
    checks.append(("Validation measures on the CSVs = pipeline/validation.py on gold", diff.empty,
                   f"{len(both)} year x indicator values compared, {len(diff)} differ"))
    districts = d["measures"].dropna(subset=["district"])
    m = districts.merge(d["mart"], on=["district", "year"], how="outer", suffixes=("", "_mart"), indicator=True)
    worst, cells = 0.0, 0
    for col, (mart_col, scale) in MART_COLUMNS.items():
        mart_values = m[mart_col if mart_col != col else f"{col}_mart"].astype(float)
        gap = (m[col].astype(float) * scale - mart_values).abs()
        worst, cells = max(worst, float(gap.max())), cells + len(gap)
    ok = (m["_merge"] == "both").all() and worst <= 0.001 + 1e-9  # the mart is rounded to 3 decimals
    checks.append(("District measures on the CSVs = mart_district_year in gold", bool(ok),
                   f"{len(m)} district-years, {cells:,} values compared, largest difference {worst:.4f} "
                   "(the mart is rounded to 3 decimals)"))
    state = d["measures"][d["measures"]["district"].isna()].set_index("year")
    summed = districts.groupby("year")[ADDITIVE].sum()
    bad = [c for c in ADDITIVE if not np.allclose(summed[c], state.loc[summed.index, c])]
    checks.append(("District rows add up to the state row (additive measures)", not bad,
                   f"{len(ADDITIVE)} measures x {len(summed)} years; differ: {', '.join(bad) or 'none'}"))
    return checks


# --- workbook -----------------------------------------------------------------------

def style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def finish(ws, widths: dict[str, int]):
    for row in ws.iter_rows():
        for cell in row:
            if cell.font.name != FONT:
                cell.font = Font(name=FONT, bold=cell.font.bold, color=cell.font.color, size=cell.font.size,
                                 italic=cell.font.italic)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def title(ws, text: str, note: str):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = note
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")


def reconciliation_sheet(wb, d: dict):
    ws = wb.active
    ws.title = "Reconciliation"
    ref = C.load_reference("official_totals.json")
    title(ws, "Reconciliation: pipeline totals vs official UDISE+ figures, Rajasthan", f"{C.SOURCE_NOTE}.")
    rows = d["pipeline"].merge(d["csv_validation"], on=["year", "indicator"], suffixes=("_gold", "_csv"))
    order = {k: i for i, k in enumerate(INDICATOR_LABELS)}
    rows = rows.sort_values(["year", "indicator"], key=lambda s: s.map(order) if s.name == "indicator" else s)
    first, last = 8, 8 + len(rows) - 1
    ws["A4"], ws["B4"] = "Exact matches", f'=COUNTIF(J{first}:J{last},"exact match")'
    ws["C4"], ws["D4"] = "of published figures", f"=COUNT(D{first}:D{last})"
    ws["A5"], ws["B5"] = "Export differs from gold", f'=COUNTIF(I{first}:I{last},"<>0")'
    for c in ("A4", "A5", "C4"):
        ws[c].font = Font(name=FONT, bold=True)
    ws["F4"] = "Blue = values written by the script (official figures, pipeline results). Black = Excel formulas."
    ws["F4"].font = Font(name=FONT, italic=True, color="595959")
    headers = ["Year", "Indicator", "Description", "Official", "Pipeline (gold DB)", "Power BI CSVs",
               "Difference (pipeline - official)", "% difference", "CSV - gold", "Result"]
    for i, h in enumerate(headers, 1):
        ws.cell(7, i, h)
    style_header(ws, 7, len(headers))
    for r, row in enumerate(rows.itertuples(index=False), first):
        official = ref["totals"].get(row.year, {}).get(row.indicator)
        ws.cell(r, 1, row.year)
        ws.cell(r, 2, row.indicator)
        ws.cell(r, 3, INDICATOR_LABELS[row.indicator])
        if official is not None:
            ws.cell(r, 4, official).comment = Comment(f"Source: {ref['source']}", "pipeline")
        ws.cell(r, 5, int(row.value_gold))
        ws.cell(r, 6, int(row.value_csv))
        ws.cell(r, 7, f'=IF(D{r}="","",E{r}-D{r})')
        ws.cell(r, 8, f'=IF(D{r}="","",G{r}/D{r})')
        ws.cell(r, 9, f"=F{r}-E{r}")
        ws.cell(r, 10, f'=IF(I{r}<>0,"export differs from gold",IF(D{r}="","no official figure",'
                       f'IF(G{r}=0,"exact match","differs")))')
        for c in (4, 5, 6):
            ws.cell(r, c).font = Font(name=FONT, color=BLUE)
            ws.cell(r, c).number_format = INDIAN
        for c in (7, 9):
            ws.cell(r, c).number_format = FORMATS["signed"]
        ws.cell(r, 8).number_format = "0.000%"
        for c in range(1, 11):
            ws.cell(r, c).border = Border(bottom=THIN)
    ws.cell(last + 2, 1, "Enrolment includes pre-primary, as published. Government = management groups marked "
                         "government in reference/management_codes.json. 2025-26 has no published figure yet.")
    ws.cell(last + 2, 1).font = Font(name=FONT, italic=True, color="595959")
    ws.freeze_panes = "A8"
    finish(ws, {"A": 26, "B": 24, "C": 36, "D": 14, "E": 18, "F": 16, "G": 18, "H": 13, "I": 12, "J": 24})


def state_sheet(wb, d: dict) -> dict[str, int]:
    """Measures as rows, years as columns. Returns measure -> row number (for the district sheet)."""
    ws = wb.create_sheet("State measures")
    title(ws, "Expected values of the DAX measures, Rajasthan (state level)",
          "Computed in SQL over the exported CSVs (sql/powerbi/check_district_measures.sql). "
          "In Power BI, a card with the measure and only the year filtered should show this value.")
    state = d["measures"][d["measures"]["district"].isna()].set_index("year")
    years = list(state.index)
    headers = ["DAX measure (MEASURES.md)", "Column in check SQL", *years, "Change", "Change %"]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, len(headers))
    rows = {}
    for r, (col, name, fmt) in enumerate(MEASURES, 5):
        rows[col] = r
        ws.cell(r, 1, name)
        ws.cell(r, 2, col)
        for j, y in enumerate(years):
            cell = ws.cell(r, 3 + j, float(state.loc[y, col]))
            cell.number_format, cell.font = FORMATS[fmt], Font(name=FONT, color=BLUE)
        a, b = get_column_letter(3), get_column_letter(2 + len(years))
        ws.cell(r, 3 + len(years), f"={b}{r}-{a}{r}").number_format = \
            "+0.0%;-0.0%;0.0%" if fmt == "pct" else ("+0.00;-0.00;0.00" if fmt == "ratio" else "+#,##0;-#,##0;0")
        if fmt != "pct":  # a % change of a percentage is misleading: the change in points is enough
            ws.cell(r, 4 + len(years), f'=IF({a}{r}=0,"",({b}{r}-{a}{r})/ABS({a}{r}))').number_format = \
                "+0.0%;-0.0%;0.0%"
    note = len(MEASURES) + 6
    ws.cell(note, 1, f"Change = {years[-1]} minus {years[0]} (the YoY measures). For % measures the change is in "
                     "percentage points. District indicators count only schools with students in classes 1-12.")
    ws.cell(note, 1).font = Font(name=FONT, italic=True, color="595959")
    ws.freeze_panes = "C5"
    finish(ws, {"A": 36, "B": 28, "C": 14, "D": 14, "E": 12, "F": 11})
    return rows


def district_sheet(wb, d: dict, state_rows: dict[str, int]):
    ws = wb.create_sheet("District measures")
    title(ws, "Expected values of the DAX measures, by district and year",
          "Columns marked 'vs state' are formulas: district value minus the state value of the same year "
          "(sheet State measures), as the 'vs State' DAX measures.")
    districts = d["measures"].dropna(subset=["district"]).sort_values(["district", "year"])
    years = list(d["measures"][d["measures"]["district"].isna()]["year"])
    vs_state = ["ptr", "pct_single_teacher", "avg_infra_index", "girls_share", "teacher_shortfall_pct"]
    headers = ["District", "Year", *[name for _, name, _ in MEASURES], *[f"{dict((c, n) for c, n, _ in MEASURES)[c]} vs state"
                                                                        for c in vs_state]]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, len(headers))
    ws.row_dimensions[4].height = 45
    col_of = {c: 3 + i for i, (c, _, _) in enumerate(MEASURES)}
    year_cols = f"'State measures'!$C$4:${get_column_letter(2 + len(years))}$4"
    for r, row in enumerate(districts.itertuples(index=False), 5):
        ws.cell(r, 1, row.district)
        ws.cell(r, 2, row.year)
        for c, _, fmt in MEASURES:
            v = getattr(row, c)
            cell = ws.cell(r, col_of[c], None if pd.isna(v) else float(v))
            cell.number_format, cell.font = FORMATS[fmt], Font(name=FONT, color=BLUE)
        for k, c in enumerate(vs_state):
            sr = state_rows[c]
            state_range = f"'State measures'!$C${sr}:${get_column_letter(2 + len(years))}${sr}"
            fmt = dict((m, f) for m, _, f in MEASURES)[c]
            cell = ws.cell(r, 3 + len(MEASURES) + k,
                           f"={get_column_letter(col_of[c])}{r}-INDEX({state_range},MATCH(B{r},{year_cols},0))")
            cell.number_format = "+0.0%;-0.0%;0.0%" if fmt == "pct" else "+0.00;-0.00;0.00"
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4 + len(districts)}"
    finish(ws, {"A": 20, "B": 9, **{get_column_letter(i): 13 for i in range(3, len(headers) + 1)}})


def checks_sheet(wb, checks):
    ws = wb.create_sheet("Checks")
    title(ws, "Automated checks run by scripts/make_reconciliation_workbook.py",
          "The script stops with exit code 1 if any check fails.")
    for i, h in enumerate(["Check", "Result", "Detail"], 1):
        ws.cell(4, i, h)
    style_header(ws, 4, 3)
    for r, (name, ok, detail) in enumerate(checks, 5):
        ws.cell(r, 1, name)
        ws.cell(r, 2, "PASS" if ok else "FAIL").font = Font(name=FONT, bold=True, color="008000" if ok else "C00000")
        ws.cell(r, 3, detail)
    finish(ws, {"A": 62, "B": 9, "C": 90})


def notes_sheet(wb, gold_db: Path):
    ws = wb.create_sheet("Notes")
    title(ws, "Definitions, norms and sources", f"{C.SOURCE_NOTE}.")
    norms = C.load_reference("norms.json")
    lines = [
        ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Gold database", C.display_path(gold_db)),
        ("Regenerate", "python scripts/export_powerbi.py, then python scripts/make_reconciliation_workbook.py"),
        ("", ""),
        ("Official", "reference/official_totals.json: " + C.load_reference("official_totals.json")["source"]),
        ("Pipeline (gold DB)", "pipeline/validation.py running sql/validation/01_official_indicators.sql on "
                               "int_school_year in the gold database"),
        ("Power BI CSVs", "sql/powerbi/check_validation_measures.sql: the validation DAX measures written in SQL "
                          "over the exported facts and dimensions (a second, independent path)"),
        ("District indicators", "Only schools with students in classes 1-12, as in mart_district_year. "
                                "PTR = students in classes 1-12 / teachers."),
        ("State value", "Ratio of state totals (all schools pooled), not the mean of the 41 district values."),
        ("", ""),
        ("Teacher norm", norms["note"]),
    ]
    lines += [(f"  {n['stage']} (classes {n['classes']})",
               f"{n['pupils_per_teacher']} pupils per teacher, minimum {n['min_teachers']}. {n['source']}. [{n['status']}]")
              for n in norms["teachers"]]
    lines += [("Teachers for norms", "total_tch minus teachers who teach only pre-primary (the norms are for classes 1-12)."),
              ("Teacher shortfall", "Sum over schools of max(0, required - teachers): a surplus in one school does "
                                    "not fill a gap in another. Net gap = required - teachers (can be negative)."),
              ("Classroom norm", f"{norms['classrooms']['rule']}. {norms['classrooms']['source']}. "
                                 f"[{norms['classrooms']['status']}]"),
              ("Classroom shortfall", "Sum over schools of max(0, teachers - classrooms). 'Usable rooms' counts only "
                                      "classrooms in good condition or needing minor repair.")]
    for r, (k, v) in enumerate(lines, 4):
        ws.cell(r, 1, k).font = Font(name=FONT, bold=True)
        ws.cell(r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    finish(ws, {"A": 30, "B": 120})


def build(gold_db: Path = C.GOLD_DB, csv_dir: Path = CSV_DIR, out: Path = OUT) -> list:
    d = compute(gold_db, csv_dir)
    checks = run_checks(d)
    wb = Workbook()
    reconciliation_sheet(wb, d)
    state_rows = state_sheet(wb, d)
    district_sheet(wb, d, state_rows)
    checks_sheet(wb, checks)
    notes_sheet(wb, gold_db)
    wb.calculation.fullCalcOnLoad = True  # openpyxl stores no results: Excel computes them on open
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return checks


def main():
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("--gold-db", type=Path, default=C.GOLD_DB, help="gold DuckDB file (read-only)")
    p.add_argument("--csv-dir", type=Path, default=CSV_DIR, help="folder written by export_powerbi.py")
    p.add_argument("--out", type=Path, default=OUT)
    a = p.parse_args()
    checks = build(a.gold_db, a.csv_dir, a.out)
    for name, ok, detail in checks:
        print(f"  {'PASS' if ok else 'FAIL'}  {name}: {detail}")
    failed = sum(not ok for _, ok, _ in checks)
    print(f"Wrote {C.display_path(a.out)}: {len(checks) - failed} of {len(checks)} checks passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
