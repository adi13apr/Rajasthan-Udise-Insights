"""Power BI export and reconciliation workbook (synthetic data only)."""
import duckdb
import pandas as pd
import pytest
from openpyxl import load_workbook

from pipeline import config as C
from pipeline.bronze import load_bronze
from pipeline.gold import build_gold
from pipeline.silver import build_silver
from scripts import export_powerbi as X
from scripts import make_reconciliation_workbook as R
from scripts.make_sample_data import make_sample


def school(key, primary=0, upper=0, teachers=0, pre_primary_only=0, classrooms=0, usable=0, pre_primary=0):
    return {"school_key": key, "year_key": 1, "enr_pre_primary": pre_primary, "enr_1_12": primary + upper,
            "enr_girls_1_12": 0, "enr_primary": primary, "enr_primary_g": 0, "enr_upper_primary": upper,
            "enr_secondary": 0, "enr_higher_secondary": 0, "enr_secondary_plus": 0, "enr_secondary_plus_g": 0,
            "teachers": teachers, "teachers_pre_primary_only": pre_primary_only, "single_teacher": int(teachers == 1),
            "classrooms": classrooms, "classrooms_usable": usable, "infra_index": 50.0,
            "has_electricity": 1.0, "has_safe_water": 1.0, "has_girls_toilet": 1.0}


def test_resource_gap_logic_on_hand_worked_schools():
    """Norms: primary 30 (min 2), upper primary 35 (min 3); one classroom per teacher.
    School 1: primary 61 -> max(2, 3) = 3; upper 36 -> max(3, 2) = 3; needs 6, has 4 - 1 pre-primary-only = 3.
    School 2: primary 20 -> 2; has 5: surplus 3. School 3 has no students in classes 1-12: left out."""
    con = duckdb.connect()
    agg = pd.DataFrame([school(1, primary=61, upper=36, teachers=4, pre_primary_only=1, classrooms=2, usable=1),
                        school(2, primary=20, teachers=5, classrooms=6, usable=6),
                        school(3, pre_primary=25, teachers=2)])
    con.register("agg_school_year", agg)
    con.register("dim_school", pd.DataFrame({"school_key": [1, 2, 3], "district_key": [1, 1, 1]}))
    con.register("dim_district", pd.DataFrame({"district_key": [1], "district_name": ["AJMER"]}))
    con.register("dim_year", pd.DataFrame({"year_key": [1], "academic_year": ["2024-25"]}))
    con.register("norms", X.norms())
    out = con.execute((R.CHECK_SQL / "check_district_measures.sql").read_text(encoding="utf-8")).df()
    state = out[out["district"].isna()].iloc[0]
    assert state["schools_with_students"] == 2
    assert state["teachers_required"] == 8
    assert state["teacher_shortfall"] == 3 and state["teacher_surplus"] == 3 and state["net_teacher_gap"] == 0
    assert state["schools_short_of_teachers"] == 1
    assert state["classroom_shortfall"] == 2            # school 1: 4 teachers, 2 rooms
    assert state["classroom_shortfall_usable"] == 3     # school 1: 4 teachers, 1 usable room


@pytest.fixture(scope="module")
def sample_gold(tmp_path_factory):
    d = tmp_path_factory.mktemp("pbi")
    make_sample(d / "raw")
    load_bronze(d / "raw", d / "bronze", run_id="r1")
    build_silver(d / "bronze", d / "silver", d / "quarantine", run_id="r1")
    build_gold(d / "silver", d / "udise.duckdb", d / "parquet", d / "exports")
    return d


def test_export_writes_every_table_without_school_codes(sample_gold):
    d = sample_gold
    manifest = X.export(d / "udise.duckdb", d / "csv")
    assert set(manifest["tables"]) == {*X.TABLES, "official_totals", "norms"}
    con = duckdb.connect(str(d / "udise.duckdb"), read_only=True)
    for t in ["dim_school", "fact_enrolment", "fact_teachers", "fact_school_facilities"]:
        assert manifest["tables"][t]["rows"] == con.execute(f"select count(*) from {t}").fetchone()[0], t
    # agg_school_year has the grain of the school-year snapshot
    assert manifest["tables"]["agg_school_year"]["rows"] == manifest["tables"]["fact_school_facilities"]["rows"]
    con.close()
    assert "pseudocode" not in pd.read_csv(d / "csv" / "dim_school.csv", nrows=1).columns
    official = pd.read_csv(d / "csv" / "official_totals.csv")
    assert len(official) == 2 * len(X.INDICATOR_LABELS)    # every year x indicator, empty where unpublished


def test_reconciliation_checks_pass_and_workbook_has_its_sheets(sample_gold):
    d = sample_gold
    X.export(d / "udise.duckdb", d / "csv")
    checks = R.build(d / "udise.duckdb", d / "csv", d / "reconciliation.xlsx")
    assert all(ok for _, ok, _ in checks), [c for c in checks if not c[1]]
    wb = load_workbook(d / "reconciliation.xlsx")
    assert wb.sheetnames == ["Reconciliation", "State measures", "District measures", "Checks", "Notes"]
    assert wb["Reconciliation"]["J8"].value.startswith("=IF(")   # results are Excel formulas


def test_norms_reference():
    n = C.load_reference("norms.json")
    stages = {t["stage"]: t for t in n["teachers"]}
    assert set(stages) == {"primary", "upper_primary", "secondary", "higher_secondary"}
    assert stages["primary"]["pupils_per_teacher"] == 30 and stages["upper_primary"]["pupils_per_teacher"] == 35


def test_powerbi_project_is_consistent(sample_gold):
    """The generated .pbip: every measure of MEASURES.md is in the model, every relationship and every
    visual field points at a real column or measure (Power BI itself is not needed for this test)."""
    import json
    import re

    from scripts import make_powerbi_project as P
    d = sample_gold
    X.export(d / "udise.duckdb", d / "csv")
    counts = P.build(d / "pbi", d / "csv")
    measures, calc = P.parse_measures()
    assert counts == {"tables": 14, "measures": len(measures), "calculated_columns": 2, "relationships": 13, "pages": 3}
    assert {c["name"] for c in calc} == {"teachers_required", "teachers_for_norms"}
    assert len(measures) == len({m["name"] for m in measures}) and all(m["folder"] for m in measures)
    model = d / "pbi" / f"{P.NAME}.SemanticModel" / "definition"
    tmdl = (model / "tables" / f"{P.MEASURE_TABLE}.tmdl").read_text(encoding="utf-8")
    assert all(f"\tmeasure {P.q(m['name'])} =" in tmdl for m in measures)
    columns = {t: {c for c, _ in cols} for t, cols in P.csv_schema(d / "csv").items()}
    columns["agg_school_year"] |= {c["name"] for c in calc}
    for many, one in P.RELATIONSHIPS:
        for ref in (many, one):
            table, col = ref.split(".")
            assert col in columns[table], ref
    names = {m["name"] for m in measures}
    visuals = list((d / "pbi" / f"{P.NAME}.Report" / "definition" / "pages").rglob("visual.json"))
    assert len(visuals) > 20
    checked = 0
    for v in visuals:
        text = v.read_text(encoding="utf-8")
        json.loads(text)
        for entity, prop in re.findall(r'"Entity": "(\w+)"\s*}\s*},\s*"Property": "([^"]+)"', text):
            checked += 1
            if entity == P.MEASURE_TABLE:
                assert prop in names, (v.parent.name, prop)
            else:
                assert prop in columns[entity], (v.parent.name, entity, prop)
    assert checked > 50   # the pattern really finds the field references
