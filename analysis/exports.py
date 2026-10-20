"""A5: district-level outputs for Tableau and Excel (aggregates only; fine to publish).

    data/gold/exports/tableau_district_indicators.csv       one row per district and year (wide)
    data/gold/exports/tableau_district_indicators_long.csv  one row per district, year and indicator
    excel/rajasthan_udise_district_summary.xlsx             summary sheets + a Data table for PivotTables
"""
import re
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from pipeline import config as C

from . import data as D

EXCEL_PATH = C.ROOT / "excel" / "rajasthan_udise_district_summary.xlsx"
INDICATORS = {
    "schools": "Schools (with students in classes 1-12)",
    "enrolment": "Students in classes 1-12",
    "teachers": "Teachers",
    "ptr": "Pupil-teacher ratio",
    "pct_single_teacher": "Single-teacher schools (%)",
    "avg_infra_index": "Infrastructure index (0-100)",
    "girls_share": "Girls' share, classes 1-12",
    "girls_share_primary": "Girls' share, classes 1-5",
    "girls_share_sec_plus": "Girls' share, classes 9-12",
    "girls_retention_gap": "Girls' retention gap (share 9-12 minus share 1-5)",
    "boys_progression": "Boys' progression ratio (9-12 vs 1-5)",
    "girls_progression": "Girls' progression ratio (9-12 vs 1-5)",
    "sc_share": "SC share of students", "st_share": "ST share of students",
    "pct_has_electricity": "Schools with electricity (%)", "pct_has_safe_water": "Schools with safe water (%)",
    "pct_has_girls_toilet": "Schools with a girls' toilet (%)",
    "pct_has_boys_toilet": "Schools with a boys' toilet (%)", "pct_has_handwash": "Schools with handwash (%)",
    "pct_has_library": "Schools with a library (%)", "pct_has_playground": "Schools with a playground (%)",
    "pct_has_ramp": "Schools with a ramp (%)", "pct_has_internet": "Schools with internet (%)",
    "pct_has_computer": "Schools with a computer (%)",
}


def tableau_frame() -> pd.DataFrame:
    dy = D.district_year()
    latest = dy[dy["year"] == D.LATEST].copy()
    latest["pattern"] = D.patterns(latest)
    matched = D.matched_single_teacher()[["district", "change_pp", "teachers_change_pct"]].rename(
        columns={"change_pp": "single_teacher_change_matched_pp",
                 "teachers_change_pct": "teachers_change_matched_pct"})
    out = dy.merge(latest[["district", "pattern"]], on="district", how="left").merge(matched, on="district",
                                                                                      how="left")
    for col, worst_high in (("avg_infra_index", False), ("pct_single_teacher", True), ("ptr", True)):
        out[f"rank_{col}"] = out.groupby("year")[col].rank(ascending=not worst_high, method="min").astype(int)
    out["state"] = "Rajasthan"  # helps Tableau's geocoding
    return out


def write_tableau(exports_dir=C.GOLD_EXPORTS_DIR) -> pd.DataFrame:
    wide = tableau_frame()
    wide.to_csv(exports_dir / "tableau_district_indicators.csv", index=False)
    long = wide.melt(id_vars=["state", "district", "year", "pattern"], value_vars=list(INDICATORS),
                     var_name="indicator", value_name="value")
    long["indicator_label"] = long["indicator"].map(INDICATORS)
    long.to_csv(exports_dir / "tableau_district_indicators_long.csv", index=False)
    return wide


# --- Excel ------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="E8EEF8")
FIXED_TIME = datetime(2026, 1, 1)  # timestamps fixed so the file only changes when the data does


def _normalise_zip(path) -> None:
    """An .xlsx is a zip; rewrite it with a fixed timestamp on every entry and in its
    properties (openpyxl stamps 'modified' with the save time)."""
    with zipfile.ZipFile(path) as z:
        entries = [(info, z.read(info)) for info in z.infolist()]
    stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    entries = [(info, re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp, data)
                if info.filename == "docProps/core.xml" else data) for info, data in entries]
    with zipfile.ZipFile(path, "w") as z:
        for info, data in entries:
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED_TIME.timetuple()[:6])
            fixed.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(fixed, data)
NOTE = ("Summary sheets are pivot tables computed in Python (openpyxl cannot create native Excel "
        "PivotTables). To explore further: select any cell in the 'Data' sheet's table, then "
        "Insert > PivotTable.")


def _write_sheet(wb: Workbook, title: str, df: pd.DataFrame, number_format: str = "#,##0.000",
                 as_table: str | None = None, intro: str = "") -> None:
    ws = wb.create_sheet(title)
    start = 1
    if intro:
        ws.cell(1, 1, intro).font = Font(italic=True, color="52514E")
        start = 3
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(start, j, str(col))
        c.font, c.fill = Font(bold=True), HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(df.itertuples(index=False), start=start + 1):
        for j, v in enumerate(row, start=1):
            v = None if (isinstance(v, float) and pd.isna(v)) else v
            c = ws.cell(i, j, v.item() if hasattr(v, "item") else v)
            if isinstance(c.value, float):
                c.number_format = number_format
            elif isinstance(c.value, int):
                c.number_format = "#,##0"
    for j, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(str(x)) for x in df.iloc[:, j - 1].head(200))) if len(df) else 10
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 40)
    ws.freeze_panes = ws.cell(start + 1, 2)
    if as_table:
        ref = f"A{start}:{get_column_letter(len(df.columns))}{start + len(df)}"
        table = Table(displayName=as_table, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)


def write_excel(path=EXCEL_PATH) -> None:
    dy = D.district_year()
    latest = dy[dy["year"] == D.LATEST].copy()
    latest["pattern"] = D.patterns(latest)
    mgmt = D.query("select * from mart_district_year_management order by district, year, management_group")

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = ["Rajasthan school education: district summary (UDISE+ 2024-25 and 2025-26)", "",
             f"{C.SOURCE_NOTE}. District-level aggregates only.",
             "Built by analysis/exports.py from the gold marts of the pipeline (python -m analysis.run).",
             "Schools with no students in classes 1-12 are excluded. Definitions: docs/data_dictionary.md.",
             "", NOTE, "", "Sheets:",
             "  Rankings 2025-26: the 41 districts with infrastructure, teachers and girls' indicators",
             "  Enrolment pivot: students in classes 1-12 by district and management group, 2025-26",
             "  Schools pivot: schools by district and year, with change",
             "  Girls patterns: retention gap and progression ratios by district, 2025-26",
             "  Data: district x year x management group (Excel Table, for your own PivotTables)"]
    for i, line in enumerate(lines, start=1):
        ws.cell(i, 1, line).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    ws.column_dimensions["A"].width = 110

    rank = latest[["district", "schools", "enrolment", "teachers", "ptr", "pct_single_teacher",
                   "avg_infra_index", "girls_share_sec_plus", "girls_retention_gap", "pattern"]] \
        .sort_values("avg_infra_index").rename(columns=INDICATORS)
    _write_sheet(wb, "Rankings 2025-26", rank, intro="Sorted by infrastructure index, lowest first.")

    pivot = pd.pivot_table(mgmt[mgmt["year"] == D.LATEST], index="district", columns="management_group",
                           values="enrolment", aggfunc="sum", margins=True, margins_name="Total").reset_index()
    _write_sheet(wb, "Enrolment pivot", pivot, number_format="#,##0",
                 intro="Students in classes 1-12, 2025-26, by management group.")

    schools = pd.pivot_table(dy, index="district", columns="year", values="schools", aggfunc="sum",
                             margins=True, margins_name="Total").reset_index()
    schools["change"] = schools[D.LATEST] - schools[D.EARLIER]
    _write_sheet(wb, "Schools pivot", schools, number_format="#,##0",
                 intro="Schools with students in classes 1-12. ALWAR/KOTPUTLI-BEHROR change partly "
                       "reflects 180 schools moving district (see docs/design_decisions.md).")

    girls = latest[["district", "pattern", "girls_share_primary", "girls_share_sec_plus", "girls_retention_gap",
                    "boys_progression", "girls_progression", "st_share"]].sort_values("girls_retention_gap") \
        .rename(columns=INDICATORS)
    _write_sheet(wb, "Girls patterns", girls, intro="Pattern groups: see notebooks/analysis.ipynb, section A2.")

    _write_sheet(wb, "Data", mgmt, as_table="tblDistrictManagement")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.properties.created = wb.properties.modified = FIXED_TIME
    wb.save(path)
    _normalise_zip(path)


def build() -> dict:
    wide = write_tableau()
    write_excel()
    return {"tableau_rows": len(wide), "excel": C.display_path(EXCEL_PATH)}
